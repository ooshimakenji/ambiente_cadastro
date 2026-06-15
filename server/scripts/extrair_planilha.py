# -*- coding: utf-8 -*-
"""
Extrator de massa de TESTE (dev) da planilha real do dono → JSON.
Lê 'PLANILHA - SERVIÇOS IMEDIATOS.xlsx' (read-only) e emite o que o importador
Prisma vai inserir no ambiental_cadastro. NÃO escreve nada na planilha.

Escopo (decidido com o usuário):
- OS magras: últimas N AS (de baixo p/ cima) da aba 'Controle de ASs em Aberto' col0
  (as mais recentes / junho) → viram OS status ABERTA, sem tipo/equipe.
- Referência: Tipos (Apoio col6) e Equipes (Apoio col0) p/ os dropdowns do app.
- Folhas Casa: PULADAS (dado muito sujo, decisão do usuário).

Uso: python scripts/extrair_planilha.py
Env: XLSX_PATH (default abaixo), IMPORT_N (default 500), IMPORT_OUT (default server/out/planilha.json)
"""
import json
import os
import re
import sys

# Console do Windows usa cp1252 por padrão e quebra ao imprimir '→'/acentos.
# Força UTF-8 no stdout/stderr (no-op em terminais que já são UTF-8).
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl não instalado. Rode: pip install openpyxl")

XLSX = os.environ.get(
    "XLSX_PATH",
    r"C:\Users\tanuk\Downloads\power_BI\PLANILHA - SERVIÇOS IMEDIATOS.xlsx",
)
N = int(os.environ.get("IMPORT_N", "500"))
AQUI = os.path.dirname(os.path.abspath(__file__))
OUT = os.environ.get("IMPORT_OUT", os.path.join(AQUI, "..", "out", "planilha.json"))

SEQ_RE = re.compile(r"20\d{8}")


def distintos_coluna(ws, col_idx, pular_header=True):
    """Valores distintos não-vazios de uma coluna (0-based), preservando 1ª ocorrência."""
    vistos, ordem = set(), []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if pular_header and i == 0:
            continue
        if col_idx >= len(row):
            continue
        c = row[col_idx]
        if c is None:
            continue
        s = str(c).strip()
        if s and s not in vistos:
            vistos.add(s)
            ordem.append(s)
    return ordem


def main():
    if not os.path.exists(XLSX):
        sys.exit(f"Planilha não encontrada: {XLSX}")
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

    # --- Apoio: tipos (col6) e equipes (col0) ---
    apoio = wb["Apoio"]
    tipos = distintos_coluna(apoio, 6)
    equipes = distintos_coluna(apoio, 0)

    # --- Controle de ASs em Aberto: col0 = sequenciais; pega as últimas N ---
    ctrl = wb["Controle de ASs em Aberto"]
    seqs = []
    for row in ctrl.iter_rows(min_col=1, max_col=1, values_only=True):
        c = row[0]
        if c is None:
            continue
        s = str(c).strip()
        if SEQ_RE.fullmatch(s):
            seqs.append(s)
    # últimas N (de baixo p/ cima = mais recentes), únicas, preservando ordem
    ultimas = seqs[-N:]
    vistos, sequenciais = set(), []
    for s in ultimas:
        if s not in vistos:
            vistos.add(s)
            sequenciais.append(s)

    dados = {
        "origem": os.path.basename(XLSX),
        "tipos": tipos,
        "equipes": equipes,
        "sequenciais_abertos": sequenciais,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)

    print(f"OK → {OUT}")
    print(f"  tipos={len(tipos)} · equipes={len(equipes)} · sequenciais={len(sequenciais)} (de {len(seqs)} no total)")


if __name__ == "__main__":
    main()
